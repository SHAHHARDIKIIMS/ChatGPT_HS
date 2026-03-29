"""Quick HTML preview of the scraped doctor data."""
from flask import Flask
from openpyxl import load_workbook

app = Flask(__name__)

@app.route("/")
def index():
    wb = load_workbook("doctors_sikarin_hatyai.xlsx")
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

    rows_html = ""
    for r in range(2, ws.max_row + 1):
        cells = ""
        for c in range(1, ws.max_column + 1):
            val = ws.cell(r, c).value or ""
            val = str(val).replace("\n", "<br>")
            if headers[c - 1] == "Photo URL" and val:
                cells += f'<td><img src="{val}" style="width:60px;height:auto;border-radius:4px"></td>'
            elif headers[c - 1] == "Profile URL" and val:
                cells += f'<td><a href="{val}" target="_blank">View</a></td>'
            else:
                cells += f"<td>{val}</td>"
        rows_html += f"<tr>{cells}</tr>"

    header_html = "".join(f"<th>{h}</th>" for h in headers)

    return f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>Sikarin Doctors</title>
<style>
body {{ font-family: sans-serif; margin: 20px; }}
h1 {{ color: #2E86C1; }}
table {{ border-collapse: collapse; width: 100%; font-size: 13px; }}
th {{ background: #2E86C1; color: white; padding: 8px; text-align: left; position: sticky; top: 0; }}
td {{ border: 1px solid #ddd; padding: 6px; vertical-align: top; max-width: 250px; word-wrap: break-word; }}
tr:nth-child(even) {{ background: #f9f9f9; }}
.count {{ color: #666; margin-bottom: 16px; }}
</style></head><body>
<h1>Sikarin Hospital Hat Yai — Doctors Directory</h1>
<p class="count">{ws.max_row - 1} doctors scraped from
<a href="https://hatyai.sikarin.com/doctors">hatyai.sikarin.com/doctors</a></p>
<table><thead><tr>{header_html}</tr></thead><tbody>{rows_html}</tbody></table>
</body></html>"""

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8080)
