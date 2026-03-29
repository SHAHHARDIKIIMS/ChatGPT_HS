"""
Scrape all doctor data from Sikarin Hospital Hat Yai website
(https://hatyai.sikarin.com/doctors) via the WordPress REST API
and export to an Excel file.
"""

import requests
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

API_BASE = "https://hatyai.sikarin.com/wp-json/wp/v2"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/120.0.0.0 Safari/537.36"
    ),
}

DAY_KEYS = ["mon", "tue", "wed", "thu", "fri", "sat", "sun"]
DAY_LABELS = {
    "mon": "Monday",
    "tue": "Tuesday",
    "wed": "Wednesday",
    "thu": "Thursday",
    "fri": "Friday",
    "sat": "Saturday",
    "sun": "Sunday",
}
OUTPUT_FILE = "doctors_sikarin_hatyai.xlsx"


def fetch_clinic_taxonomy() -> dict[int, str]:
    """Fetch clinics_centers taxonomy and return {id: name} mapping."""
    url = f"{API_BASE}/clinics_centers"
    params = {"per_page": 100}
    resp = requests.get(url, headers=HEADERS, params=params, timeout=30)
    resp.raise_for_status()
    return {item["id"]: item["name"] for item in resp.json()}


def fetch_all_doctors() -> list[dict]:
    """Fetch all doctor posts via WP REST API (handles pagination)."""
    doctors = []
    page = 1
    while True:
        url = f"{API_BASE}/doctor"
        params = {"per_page": 100, "page": page}
        resp = requests.get(url, headers=HEADERS, params=params, timeout=30)
        resp.raise_for_status()
        batch = resp.json()
        if not batch:
            break
        doctors.extend(batch)
        total_pages = int(resp.headers.get("X-WP-TotalPages", 1))
        if page >= total_pages:
            break
        page += 1
    return doctors


def clean_text(text: str | None) -> str:
    """Strip whitespace, collapse newlines, and return clean text."""
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def format_availability(acf: dict) -> str:
    """Build a human-readable availability string from day fields."""
    lines = []
    for key in DAY_KEYS:
        day_data = acf.get(key, {})
        start = (day_data.get("start_time") or "").strip()
        end = (day_data.get("end_time") or "").strip()
        label = DAY_LABELS[key]
        if start and end:
            lines.append(f"{label}: {start} - {end}")
        else:
            lines.append(f"{label}: Off")
    return "\n".join(lines)


def build_row(doctor: dict, clinic_map: dict[int, str]) -> dict:
    """Transform a raw doctor API record into a flat row dict."""
    acf = doctor.get("acf", {})
    title = doctor.get("title", {}).get("rendered", "").strip()

    clinic_ids = doctor.get("clinics_centers", [])
    clinics = ", ".join(
        clinic_map.get(cid, str(cid)) for cid in clinic_ids
    )

    return {
        "Name": title,
        "Position": clean_text(acf.get("position")),
        "Clinic / Center": clinics,
        "Education / Diploma": clean_text(acf.get("doctor_info")),
        "Experience": clean_text(acf.get("experience")),
        "Association": clean_text(acf.get("association")),
        "Languages": clean_text(acf.get("lang")),
        "Availability": format_availability(acf),
        "Photo URL": acf.get("doctor_image") or "",
        "Profile URL": doctor.get("link", ""),
        "Remark": clean_text(acf.get("remark")),
    }


def write_excel(rows: list[dict], path: str) -> None:
    """Write rows to a styled Excel workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Doctors"

    columns = [
        "Name",
        "Position",
        "Clinic / Center",
        "Education / Diploma",
        "Experience",
        "Association",
        "Languages",
        "Availability",
        "Photo URL",
        "Profile URL",
        "Remark",
    ]

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2E86C1", end_color="2E86C1", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_alignment = Alignment(vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_idx, row_data in enumerate(rows, start=2):
        for col_idx, col_name in enumerate(columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row_data.get(col_name, ""))
            cell.alignment = cell_alignment
            cell.border = thin_border

    col_widths = {
        "Name": 30,
        "Position": 25,
        "Clinic / Center": 30,
        "Education / Diploma": 45,
        "Experience": 45,
        "Association": 25,
        "Languages": 20,
        "Availability": 35,
        "Photo URL": 40,
        "Profile URL": 50,
        "Remark": 25,
    }
    for col_idx, col_name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = col_widths.get(col_name, 20)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    wb.save(path)


def main() -> None:
    print("Fetching clinic/center taxonomy...")
    clinic_map = fetch_clinic_taxonomy()
    print(f"  Found {len(clinic_map)} clinics/centers")

    print("Fetching all doctors...")
    doctors = fetch_all_doctors()
    print(f"  Found {len(doctors)} doctors")

    print("Processing doctor records...")
    rows = [build_row(d, clinic_map) for d in doctors]
    rows.sort(key=lambda r: r["Name"])

    print(f"Writing {len(rows)} records to {OUTPUT_FILE}...")
    write_excel(rows, OUTPUT_FILE)
    print(f"Done! File saved: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
